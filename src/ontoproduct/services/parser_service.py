from collections import Counter
from collections.abc import Callable
from pathlib import Path

import openpyxl
import pdfplumber
from pdfplumber.page import Page
from pdfplumber.table import Table

from ontoproduct.schemas.product import FileReference, ParsedDocument
from ontoproduct.services.document_service import DocumentService

# (page, text) 한 쌍. TXT·XLSX는 page가 None, PDF는 1부터 시작하는 페이지 번호.
ParsedPage = tuple[int | None, str]

# 칸 구분선이 없는 PDF 표에서 이 간격(pt)보다 넓게 떨어진 단어는 다른 칸으로 본다.
COLUMN_GAP = 15


class ParserService:
    """업로드된 TXT/PDF/XLSX 파일을 ParsedDocument dictionary 목록으로 읽는다.

    ParsedDocument에 시트·셀 필드가 없으므로 Excel 위치는 text 안에 남긴다.
    텍스트가 없는 스캔 PDF는 OCR을 지원하지 않으므로 오류로 처리한다.
    """

    def __init__(self, workspace: str | Path) -> None:
        self.documents = DocumentService(workspace)

    def parse(self, references: list[dict]) -> list[dict]:
        """업로드 폴더 안의 원본 파일을 열어 파일·페이지별 원문 목록을 반환한다."""
        # 조작된 path로 업로드 폴더 밖의 파일을 읽지 않도록 업로드 때와 같은 검증을 거친다.
        self.documents.validate_references(references)
        refs = [FileReference.model_validate(value) for value in references]
        name_counts = Counter(ref.name for ref in refs)
        parsed = []
        for ref in refs:
            reader = READERS.get(Path(ref.name).suffix.lower())
            if reader is None:
                raise ValueError(f"{ref.name}: 지원하지 않는 파일 형식입니다")
            # 같은 이름의 파일이 여러 개면 file_id 앞 8자리로 구분한다.
            source = (
                ref.name
                if name_counts[ref.name] == 1
                else f"{ref.name} ({ref.file_id[:8]})"
            )
            # 읽기 함수가 직접 낸 오류는 메시지를 그대로 살리고, 라이브러리 오류는
            # "읽을 수 없음"으로 감싼다. 어느 쪽이든 파일명을 붙여 오류 화면에서
            # 원인을 알 수 있게 한다.
            try:
                pages = reader(Path(ref.path))
            except ValueError as exc:
                raise ValueError(f"{ref.name}: {exc}") from exc
            except Exception as exc:
                reason = f"{type(exc).__name__}: {exc}"
                raise ValueError(
                    f"{ref.name}: 문서를 읽을 수 없습니다 ({reason})"
                ) from exc
            for page, text in pages:
                document = ParsedDocument(source_file=source, text=text, page=page)
                parsed.append(document.model_dump(mode="json"))
        return parsed


def read_txt(path: Path) -> list[ParsedPage]:
    """TXT 원문을 UTF-8 또는 CP949로 읽는다."""
    data = path.read_bytes()
    # 한글 Windows 메모장은 CP949나 BOM이 붙은 UTF-8로 저장하므로 둘 다 시도한다.
    # utf-8-sig는 BOM이 없는 일반 UTF-8도 읽는다.
    for encoding in ("utf-8-sig", "cp949"):
        try:
            # Windows 줄바꿈(\r\n)을 \n으로 통일해 Extraction이 한 가지 형식만 다루게 한다.
            text = data.decode(encoding).replace("\r\n", "\n")
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ValueError("텍스트 인코딩을 판별할 수 없습니다 (UTF-8, CP949 지원)")
    if not text.strip():
        raise ValueError("읽을 수 있는 텍스트가 없습니다")
    return [(None, text)]


def read_pdf(path: Path) -> list[ParsedPage]:
    """PDF를 페이지별로 읽고, 표는 '칸 | 칸' 행으로 덧붙인다."""
    pages = []
    with pdfplumber.open(path) as pdf:
        for number, page in enumerate(pdf.pages, start=1):
            # extract_text는 표의 칸 구분을 잃으므로 표를 행 단위로 한 번 더 붙인다.
            blocks = [page.extract_text() or ""]
            for index, table in enumerate(page.find_tables(), start=1):
                cells = table.extract()
                # 세로줄 없이 가로줄만 있는 표는 행 전체가 한 칸으로 읽히므로 글자 간격으로 나눈다.
                if max(len(row) for row in cells) < 2:
                    cells = split_columns_by_gap(page, table)
                rows = [
                    " | ".join((cell or "").strip() for cell in row) for row in cells
                ]
                blocks.append(f"[Table {index}]\n" + "\n".join(rows))
            text = "\n\n".join(block for block in blocks if block.strip())
            # 빈 페이지는 건너뛰지만 page에는 원본의 실제 페이지 번호를 그대로 쓴다.
            if text:
                pages.append((number, text))
    if not pages:
        raise ValueError("읽을 수 있는 텍스트가 없습니다 (스캔 PDF OCR 미지원)")
    return pages


def split_columns_by_gap(page: Page, table: Table) -> list[list[str]]:
    """칸 구분선이 없는 표의 각 행을 단어 사이의 넓은 간격 기준으로 칸에 나눈다."""
    word_rows = [page.crop(row.bbox).extract_words() for row in table.rows]
    # 모든 행에서 넓은 간격 뒤에 시작하는 단어의 x좌표를 모아 열의 시작 위치로 삼는다.
    # 같은 칸 안의 띄어쓰기(약 3pt)는 COLUMN_GAP보다 좁아 같은 칸으로 묶인다.
    starts = []
    for words in word_rows:
        for previous, word in zip([None, *words], words, strict=False):
            if previous is None or word["x0"] - previous["x1"] > COLUMN_GAP:
                starts.append(word["x0"])
    columns: list[float] = []
    for x in sorted(starts):
        if not columns or x - columns[-1] > COLUMN_GAP:
            columns.append(x)
    # 단어를 시작 위치가 가장 가까운 왼쪽 열에 넣어, 빈 칸이 있어도 열이 밀리지 않게 한다.
    rows = []
    for words in word_rows:
        cells: list[list[str]] = [[] for _ in columns]
        for word in words:
            column = max(
                i for i, start in enumerate(columns) if start <= word["x0"] + 1
            )
            cells[column].append(word["text"])
        rows.append([" ".join(cell) for cell in cells])
    return rows


def read_xlsx(path: Path) -> list[ParsedPage]:
    """XLSX를 시트별로 읽고, 각 행을 '[Sheet: 이름, Row: n] A1=값 | B1=값' 형식으로 남긴다."""
    # data_only=True는 수식 대신 Excel이 저장해 둔 계산값을 읽는다. Excel로 한 번도 저장하지
    # 않은 파일(프로그램으로만 만든 파일)은 계산값이 없어 수식 칸이 빈 칸으로 나온다.
    # read_only=True는 큰 파일도 메모리를 적게 쓰며 읽는다.
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheets = []
        for sheet in workbook.worksheets:
            lines = []
            for row in sheet.iter_rows():
                cells = [
                    cell
                    for cell in row
                    if cell.value is not None and str(cell.value).strip()
                ]
                if cells:
                    values = " | ".join(
                        f"{cell.coordinate}={cell.value}" for cell in cells
                    )
                    lines.append(
                        f"[Sheet: {sheet.title}, Row: {cells[0].row}] {values}"
                    )
            if lines:
                sheets.append((None, "\n".join(lines)))
    finally:
        workbook.close()
    if not sheets:
        raise ValueError("시트에 읽을 수 있는 값이 없습니다")
    return sheets


READERS: dict[str, Callable[[Path], list[ParsedPage]]] = {
    ".txt": read_txt,
    ".pdf": read_pdf,
    ".xlsx": read_xlsx,
}
