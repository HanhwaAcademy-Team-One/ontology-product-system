"""Parser 테스트용 샘플 문서를 다시 만든다: uv run python tests/fixtures/make_documents.py"""
from datetime import UTC, datetime
from pathlib import Path

from fpdf import FPDF
from openpyxl import Workbook

OUT = Path(__file__).parent / "documents"
FIXED_TIME = datetime(2026, 1, 1, tzinfo=UTC)
MOTOR_TEXT = ("Product: DM-600\nClass: BLDCMotor\nManufacturer: XYZ Motors\n"
              "Rated Voltage: 24 V\nRated Power: 0.6 kW\nRated Speed: 3200 rpm\n"
              "비고: 한글 English 혼합 문서 ±5% Ω\n")


def new_pdf() -> FPDF:
    """생성 시각을 고정한 PDF 객체를 만든다."""
    pdf = FPDF()
    pdf.set_creation_date(FIXED_TIME)
    pdf.set_font("Helvetica", size=12)
    return pdf


def make_text() -> None:
    (OUT / "motor_spec.txt").write_text(MOTOR_TEXT, encoding="utf-8")
    (OUT / "motor_spec_cp949.txt").write_bytes(MOTOR_TEXT.replace("±5% Ω", "").encode("cp949"))


def make_pdf() -> None:
    pdf = new_pdf()
    pdf.add_page()
    pdf.multi_cell(0, 8, "Product: DM-600\nClass: BLDCMotor\nManufacturer: XYZ Motors")
    pdf.add_page()
    pdf.cell(0, 10, "Electrical Ratings", new_x="LMARGIN", new_y="NEXT")
    with pdf.table(col_widths=(60, 40, 30)) as table:
        for row in [("Item", "Value", "Unit"), ("Power", "0.6", "kW"),
                    ("Voltage", "24", "V"), ("Speed", "3200", "rpm")]:
            cells = table.row()
            for value in row:
                cells.cell(value)
    pdf.output(OUT / "two_page_spec.pdf")


def make_horizontal_lines_pdf() -> None:
    # 실제 사양서처럼 행마다 전체 너비의 배경 사각형과 가로줄만 있고 칸 구분선은 없는 표.
    # pdfplumber는 이런 표를 행 전체가 한 칸인 표로 인식한다.
    pdf = new_pdf()
    pdf.add_page()
    pdf.set_fill_color(235, 240, 245)
    rows = [("Item", "Value", "Unit"), ("Product", "DM-600", ""),
            ("Rated Power", "0.6", "kW"), ("Rated Speed", "3200", "rpm")]
    for index, row in enumerate(rows):
        top = 30 + index * 12
        pdf.rect(20, top, 160, 12, style="F")
        pdf.line(20, top + 12, 180, top + 12)
        for x, value in zip((22, 100, 150), row, strict=True):
            if value:
                pdf.text(x, top + 8, value)
    pdf.output(OUT / "horizontal_lines_table.pdf")


def make_scanned_pdf() -> None:
    # 도형만 있고 추출할 텍스트가 없어 스캔 페이지처럼 동작한다.
    pdf = new_pdf()
    pdf.add_page()
    pdf.rect(20, 20, 120, 40, style="F")
    pdf.output(OUT / "scanned.pdf")


def make_xlsx() -> None:
    workbook = Workbook()
    spec = workbook.active
    spec.title = "Spec"
    for row in [("Item", "Value", "Unit"), ("Product", "DM-600", None),
                ("Manufacturer", "XYZ Motors", None), ("정격 전압", 24, "V")]:
        spec.append(row)
    electrical = workbook.create_sheet("Electrical")
    electrical["A1"], electrical["B1"], electrical["C1"] = "Rated Power", 0.6, "kW"
    electrical["A3"], electrical["C3"] = "Rated Speed", "rpm"
    electrical["B3"] = 3200
    workbook.properties.created = workbook.properties.modified = FIXED_TIME.replace(tzinfo=None)
    workbook.save(OUT / "two_sheet_spec.xlsx")


def make_corrupt() -> None:
    # 정상 파일의 앞 200바이트만 남겨 손상 파일을 만든다.
    (OUT / "corrupt.pdf").write_bytes((OUT / "two_page_spec.pdf").read_bytes()[:200])
    (OUT / "corrupt.xlsx").write_bytes((OUT / "two_sheet_spec.xlsx").read_bytes()[:200])


if __name__ == "__main__":
    OUT.mkdir(exist_ok=True)
    make_text()
    make_pdf()
    make_horizontal_lines_pdf()
    make_scanned_pdf()
    make_xlsx()
    make_corrupt()
